#!/usr/bin/env python3
"""
Cost of Cultivation (COC) Report Consolidator
==============================================
Consolidates Sri Lanka Department of Agriculture "Cost of Cultivation" (COC)
Excel reports across 19 seasons (2014-2025) into a unified master training CSV.

Module: SO1 - Predictive Cost Estimator
Project: IT4010 - Intelligent Decision Support System for Agricultural Production Cost Prediction
"""

import os
import re
import sys
from pathlib import Path
import openpyxl
import pandas as pd
import numpy as np

# ==============================================================================
# 1. NORMALIZATION & ALIAS TABLES
# ==============================================================================

DISTRICT_ALIASES = {
    # Anuradhapura
    "anuradhapura": "Anuradhapura",
    "a'pura": "Anuradhapura",
    "apura": "Anuradhapura",
    "anu": "Anuradhapura",
    "anuradapura": "Anuradhapura",
    # Ampara & sub-divisions
    "ampara": "Ampara",
    "ampara east": "Ampara (East)",
    "ampara - east": "Ampara (East)",
    "ampara (east)": "Ampara (East)",
    "am east": "Ampara (East)",
    "am-east": "Ampara (East)",
    "amp east": "Ampara (East)",
    "ampara-east": "Ampara (East)",
    "amp-east": "Ampara (East)",
    "ampara west": "Ampara (West)",
    "ampara - west": "Ampara (West)",
    "ampara (west)": "Ampara (West)",
    "am west": "Ampara (West)",
    "am-west": "Ampara (West)",
    "amp west": "Ampara (West)",
    "ampara-west": "Ampara (West)",
    "amp-west": "Ampara (West)",
    "ampara wset": "Ampara (West)",
    "amara -west": "Ampara (West)",
    # Badulla
    "badulla": "Badulla",
    "bdl": "Badulla",
    "bd": "Badulla",
    # Batticaloa
    "batticaloa": "Batticaloa",
    "batti": "Batticaloa",
    # Gampaha
    "gampaha": "Gampaha",
    "gamph": "Gampaha",
    # Hambantota
    "hambantota": "Hambantota",
    "h'tota": "Hambantota",
    "h tota": "Hambantota",
    "htota": "Hambantota",
    "hamb": "Hambantota",
    # Jaffna
    "jaffna": "Jaffna",
    # Kalutara
    "kalutara": "Kalutara",
    # Kandy
    "kandy": "Kandy",
    # Kurunegala
    "kurunegala": "Kurunegala",
    "kurunagala": "Kurunegala",
    "kuru": "Kurunegala",
    "k'gala": "Kurunegala",
    # Kegalle
    "kg": "Kegalle",
    "kegalle": "Kegalle",
    # Mannar
    "mannar": "Mannar",
    # Matale
    "matale": "Matale",
    "mtl": "Matale",
    # Matara
    "matara": "Matara",
    # Monaragala
    "monaragala": "Monaragala",
    "moneragala": "Monaragala",
    "m'gala": "Monaragala",
    # Nuwara Eliya
    "nuwara eliya": "Nuwara Eliya",
    "nuwara-eliya": "Nuwara Eliya",
    "nuwara- eliya": "Nuwara Eliya",
    "nuwaraeliya": "Nuwara Eliya",
    "n'eliya": "Nuwara Eliya",
    "n.eliya": "Nuwara Eliya",
    "n.éliya": "Nuwara Eliya",
    "néliya": "Nuwara Eliya",
    "ne": "Nuwara Eliya",
    "neliya": "Nuwara Eliya",
    # Polonnaruwa
    "polonnaruwa": "Polonnaruwa",
    "polonnruwa": "Polonnaruwa",
    "polo": "Polonnaruwa",
    # Puttalam
    "puttalam": "Puttalam",
    "puttlam": "Puttalam",
    "puttalum": "Puttalam",
    "put": "Puttalam",
    # Trincomalee
    "trincomalee": "Trincomalee",
    "trinco": "Trincomalee",
    # Vavuniya
    "vavuniya": "Vavuniya",
    "vavuniy": "Vavuniya",
    # Udawalawa
    "udawalawa": "Udawalawa",
    "udawalawe": "Udawalawa",
    # Mahaweli Systems
    "system b": "Mahaweli System B",
    "sys b": "Mahaweli System B",
    "mahaweli b": "Mahaweli System B",
    "system c": "Mahaweli System C",
    "sys c": "Mahaweli System C",
    "mahaweli c": "Mahaweli System C",
    "system h": "Mahaweli System H",
    "sys h": "Mahaweli System H",
    "mahaweli h": "Mahaweli System H",
    # Whole Island aggregates
    "whole island": "Whole Island",
    "whole islaand": "Whole Island",
    "wlole island": "Whole Island",
    "wi": "Whole Island",
}

CROP_ALIASES = {
    # Paddy
    "paddy": "Paddy",
    # Maize
    "maize": "Maize",
    "maaize": "Maize",
    # Pulses
    "black gram": "Black Gram",
    "blackgram": "Black Gram",
    "black grame": "Black Gram",
    "b'gram": "Black Gram",
    "b gram": "Black Gram",
    "b'ggram": "Black Gram",
    "black g": "Black Gram",
    "cowpea": "Cowpea",
    "green gram": "Green Gram",
    "green grame": "Green Gram",
    "greengram": "Green Gram",
    # Millets & Grains
    "kurakkan": "Kurakkan",
    "finger millet": "Kurakkan",
    # Oilseeds
    "ground nut": "Groundnut",
    "groundnut": "Groundnut",
    "g'nut": "Groundnut",
    "g nut": "Groundnut",
    "gingerlly": "Gingelly",
    "gingelly": "Gingelly",
    "gingerly": "Gingelly",
    "gingelly/sesame": "Gingelly",
    "sesame": "Gingelly",
    "soybean": "Soybean",
    "soya": "Soybean",
    "soya bean": "Soybean",
    "soya beans": "Soybean",
    # Condiments & Alliums
    "big onion": "Big Onion",
    "red onion": "Red Onion",
    "red onio": "Red Onion",
    "chilli": "Green Chilli",
    "g chilli": "Green Chilli",
    "green chilli": "Green Chilli",
    "greenchilli": "Green Chilli",
    "capsicum": "Capsicum",
    "capsi": "Capsicum",
    # Vegetables
    "cabbage": "Cabbage",
    "carrot": "Carrot",
    "carro": "Carrot",
    "tomato": "Tomato",
    "tommato": "Tomato",
    "brinjal": "Brinjal",
    "brijal": "Brinjal",
    "brinal": "Brinjal",
    "bitter gourd": "Bitter Gourd",
    "bitter goard": "Bitter Gourd",
    "pole bean": "Pole Bean",
    "polebean": "Pole Bean",
    "pole been": "Pole Bean",
    "pole ben": "Pole Bean",
    "pole beans": "Pole Bean",
    "poleb": "Pole Bean",
    "pole b": "Pole Bean",
    "potato": "Potato",
    "poteto": "Potato",
    "sweet potato": "Sweet Potato",
    "manioc": "Manioc",
    "pumpkin": "Pumpkin",
    "pumking": "Pumpkin",
    "pumkin": "Pumpkin",
    "pupkin": "Pumpkin",
    "beet - root": "Beetroot",
    "beet root": "Beetroot",
    "beetroot": "Beetroot",
}

# ==============================================================================
# 2. HELPER FUNCTIONS
# ==============================================================================

def clean_str(s):
    """Clean string from whitespace and anomalous unicode characters."""
    if s is None:
        return ""
    return str(s).replace('\xa0', ' ').replace('\u2013', '-').replace('\u2014', '-').replace('\ufffd', 'e').strip()

def clean_num(val, line_idx=0):
    """
    Robust float conversion:
    - Handles commas, stray backticks, apostrophes, double quotes, asterisks.
    - Handles parenthesized numbers like '(123.45)' or '`(1579.00)'.
    - If multiline cell (e.g. '218,748\\n77,071'), extracts requested line_idx.
    - Returns np.nan on missing/non-numeric values.
    """
    if val is None:
        return np.nan
    if isinstance(val, (int, float)):
        return np.nan if np.isnan(val) else float(val)
    s = str(val).strip()
    if not s or s in ['-', '`', '*', 'na', 'nan', 'nil', '.', '/', '\\']:
        return np.nan
    if '\n' in s:
        lines = [ln.strip() for ln in s.split('\n') if ln.strip()]
        if 0 <= line_idx < len(lines):
            s = lines[line_idx]
        elif lines:
            s = lines[0]
        else:
            return np.nan
    s = re.sub(r'[\,\`\'\"\*\s]', '', s)
    m = re.match(r'^\(?([\d\.\-]+)\)?$', s)
    if m:
        try:
            return float(m.group(1))
        except (ValueError, TypeError):
            return np.nan
    return np.nan

def normalize_district(raw_text):
    """Match district from text using alias table."""
    if not raw_text:
        return None
    t = clean_str(raw_text).lower()
    t = re.sub(r'\s*-\s*', '-', t)
    t = re.sub(r'\bdistricts?\b', ' ', t)
    for k in sorted(DISTRICT_ALIASES.keys(), key=len, reverse=True):
        pattern = r'(?<![a-zA-Z0-9])' + re.escape(k) + r'(?![a-zA-Z0-9])'
        if re.search(pattern, t):
            return DISTRICT_ALIASES[k]
    return None

def normalize_crop(raw_text):
    """Match crop from text using alias table."""
    if not raw_text:
        return None
    t = clean_str(raw_text).lower()
    for k in sorted(CROP_ALIASES.keys(), key=len, reverse=True):
        pattern = r'(?<![a-zA-Z0-9])' + re.escape(k) + r'(?![a-zA-Z0-9])'
        if re.search(pattern, t):
            return CROP_ALIASES[k]
    return None

def normalize_irrigation(raw_text):
    """Extract irrigation regime if present."""
    if not raw_text:
        return ""
    t = clean_str(raw_text).lower()
    if re.search(r'\b(rain[\s-]?fed|rf)\b', t):
        return "Rainfed"
    if re.search(r'\b(irrigated|irr|irri)\b', t):
        return "Irrigated"
    return ""

def parse_season_folder(folder_name):
    """
    Parse Season, Year_Start, Year_End, Year_Range from folder name.
    """
    fn = folder_name.strip()
    season = "Maha" if "maha" in fn.lower() else "Yala"
    
    m_range = re.search(r'(\d{4})\s*[-_]\s*(\d{2,4})', fn)
    if m_range:
        y_start = int(m_range.group(1))
        y2 = m_range.group(2)
        if len(y2) == 2:
            y_end = int(str(y_start)[:2] + y2)
        else:
            y_end = int(y2)
        yr_range = f"{y_start}-{str(y_end)[-2:]}"
        return {
            "Season": season,
            "Year_Start": y_start,
            "Year_End": y_end,
            "Year_Range": yr_range
        }
    
    m_single = re.search(r'(\d{4})', fn)
    if m_single:
        y = int(m_single.group(1))
        return {
            "Season": season,
            "Year_Start": y,
            "Year_End": y,
            "Year_Range": str(y)
        }
        
    return {
        "Season": season,
        "Year_Start": np.nan,
        "Year_End": np.nan,
        "Year_Range": fn
    }

# ==============================================================================
# 3. SHEET CLASSIFICATION
# ==============================================================================

def classify_sheet(sheet, sname):
    """
    Classify sheet into:
    - 'BLANK'
    - 'COVER_TOC'
    - 'SUMMARY'
    - 'DRAFT'
    - 'CONTINUATION'
    - 'COC_TABLE'
    """
    sname_clean = clean_str(sname).strip()
    
    # Check for known draft sheets
    if sname_clean == "s":
        return "DRAFT", "Working draft copy of Whole Island; superseded by official sheet WI(Irrigated)"
        
    rows_text = []
    for r in sheet.iter_rows(max_row=15, max_col=10, values_only=True):
        txts = [str(c).strip() for c in r if c is not None and str(c).strip()]
        if txts:
            rows_text.append(" ".join(txts))
            
    if not rows_text:
        return "BLANK", "Empty worksheet"
        
    full_text = " ".join(rows_text).lower()
    sname_lower = sname.lower()
    
    # 1. Summary / statistical tables / graphs
    if any(k in sname_lower for k in ['graphs', 'stat table', 'annexes', 'summery', 'summary']) or \
       any(k in full_text for k in ['statistical table', 'land holding size', 'figure 2.1 labour', 'table 01', 'table 1 cost for land preparation']):
        return "SUMMARY", "Summary / statistical annex / graph sheet"
        
    # 2. Table of contents / Cover pages / Section dividers
    if "table of content" in full_text or "department of agriculture" in full_text or \
       "socio economics & planning centre" in full_text or "other field crops & vegetables" in full_text or \
       "page no" in full_text:
        return "COVER_TOC", "Cover page, Section Divider, or Table of Contents"
        
    has_op = any("operation" in l.lower() for l in rows_text)
    has_yield = any("yield and return" in l.lower() for l in rows_text)
    
    # 3. Multi-page continuation sheet (Yield block only)
    if not has_op and (has_yield or "cont" in full_text):
        return "CONTINUATION", "Continuation sheet for multi-page crop table"
        
    # 4. Valid COC Table MUST have an Operation table
    if has_op:
        return "COC_TABLE", "Valid Cost of Cultivation table"
        
    return "BLANK", "No Operation table found in sheet"

# ==============================================================================
# 4. METADATA PARSER (Crop, District, Irrigation)
# ==============================================================================

def parse_sheet_metadata(wb_path, sheet, sname):
    """Extract and normalize Crop, District, and Irrigation Type."""
    s_lower = clean_str(sname).lower().strip()
    
    rows_text = []
    for r in range(1, 12):
        rvals = [clean_str(sheet.cell(r, c).value) for c in range(1, 12)]
        non_empty = [v for v in rvals if v]
        if non_empty:
            rows_text.append(" ".join(non_empty))
            
    header_block = " ".join(rows_text).lower()
    header_block_norm = re.sub(r'\s*-\s*', '-', header_block)
    
    # --- 1. District Identification ---
    district = None
    
    # High Priority: Specific sub-district overrides based on Sheet Name
    # (Prevents unedited duplicated headers from overriding true sheet identity)
    if "kandy" in s_lower:
        district = "Kandy"
    elif any(k in s_lower for k in ["neliya", "n'eliya", "n.eliya", "néliya"]):
        district = "Nuwara Eliya"
    elif any(k in s_lower for k in ["sys c", "system c", "mahaweli c"]):
        district = "Mahaweli System C"
    elif any(k in s_lower for k in ["sys b", "system b", "mahaweli b"]):
        district = "Mahaweli System B"
    elif any(k in s_lower for k in ["sys h", "system h", "mahaweli h"]):
        district = "Mahaweli System H"
    elif any(k in s_lower for k in ["ampara e", "amp east", "am east", "am-east", "ampara-east"]):
        district = "Ampara (East)"
    elif any(k in s_lower for k in ["ampara w", "amp west", "am west", "am-west", "ampara-west", "ampara wset", "amara -west"]):
        district = "Ampara (West)"
        
    # Check title / header lines with sub-district specifiers
    if not district:
        if "ampara" in header_block_norm and "east" in header_block_norm:
            district = "Ampara (East)"
        elif "ampara" in header_block_norm and "west" in header_block_norm:
            district = "Ampara (West)"
        elif "system b" in header_block_norm or "sys b" in header_block_norm:
            district = "Mahaweli System B"
        elif "system c" in header_block_norm or "sys c" in header_block_norm or "mahaweli c" in header_block_norm:
            district = "Mahaweli System C"
        elif "system h" in header_block_norm or "sys h" in header_block_norm or "mahaweli h" in header_block_norm:
            district = "Mahaweli System H"

    # Check explicit "DISTRICT:" line
    if not district:
        for line in rows_text:
            m = re.search(r'district\s*[\:\-]\s*([^\d\n]+)', line, re.IGNORECASE)
            if m:
                d_cand = normalize_district(m.group(1))
                if d_cand:
                    district = d_cand
                    break
                    
    # Check top rows
    if not district:
        for line in rows_text[:5]:
            d_cand = normalize_district(line)
            if d_cand:
                district = d_cand
                break
                
    # Fallback to sheet name
    if not district:
        district = normalize_district(sname)
        
    # --- 2. Crop Identification ---
    crop = None
    for line in rows_text:
        m = re.search(r'crop\s*[\:\-]\s*([^\d\n]+)', line, re.IGNORECASE)
        if m:
            c_cand = normalize_crop(m.group(1))
            if c_cand:
                crop = c_cand
                break
    if not crop:
        for line in rows_text[:5]:
            c_cand = normalize_crop(line)
            if c_cand:
                crop = c_cand
                break
    if not crop:
        crop = normalize_crop(sname)
    if not crop and "paddy" in str(wb_path).lower():
        crop = "Paddy"
        
    # --- 3. Irrigation Regime ---
    irrigation = ""
    for line in rows_text[:5]:
        irr = normalize_irrigation(line)
        if irr:
            irrigation = irr
            break
    if not irrigation:
        irrigation = normalize_irrigation(sname)
        
    return crop, district, irrigation

# ==============================================================================
# 5. OPERATIONS TABLE PARSER
# ==============================================================================

def parse_operation_table(sheet, max_scan_row=45):
    """
    Extract Operation table subtotals:
    Labour_Cost, Machinery_Cost, Material_Cost,
    Total_Cost_Incl_Imputed, Total_Cost_Excl_Imputed
    and fallback labour days/wage rate if present.
    """
    op_row = None
    total_incl_row = None
    total_excl_row = None
    
    for r in range(1, max_scan_row + 1):
        row_str = " ".join([str(sheet.cell(r, c).value or "").strip().lower() for c in range(1, 15)])
        if "operation" in row_str and op_row is None:
            op_row = r
        if re.search(r'total\s*incl', row_str) and total_incl_row is None:
            total_incl_row = r
        if re.search(r'total\s*excl', row_str) and total_excl_row is None:
            total_excl_row = r
            
    if not op_row or not total_incl_row:
        return {}, op_row, total_incl_row, total_excl_row
        
    col_headers = {}
    for c in range(1, 16):
        txts = []
        for r in range(op_row, min(op_row + 3, total_incl_row)):
            val = sheet.cell(r, c).value
            if val is not None:
                txts.append(str(val).strip().lower())
        col_headers[c] = " ".join(txts)
        
    labour_col = None
    machinery_col = None
    material_col = None
    total_col = None
    wage_rate_col = None
    tot_lab_days_col = None
    fam_lab_days_col = None
    
    for c, h in col_headers.items():
        if "material" in h and "cost" in h:
            material_col = c
        elif "material" in h and material_col is None:
            material_col = c
            
        if ("machin" in h or "equip" in h) and machinery_col is None:
            machinery_col = c
            
        if "total" in h and ("cost" in h or "rs" in h):
            total_col = c
        elif "total" in h and total_col is None and "lab" not in h:
            total_col = c
            
        if "wage" in h or "rate" in h:
            wage_rate_col = c
        elif "fam" in h:
            fam_lab_days_col = c
        elif "total" in h and "lab" in h:
            tot_lab_days_col = c

    for c, h in col_headers.items():
        if ("labour" in h or "labor" in h) and "cost" in h:
            labour_col = c
        elif ("labour" in h or "labor" in h) and labour_col is None and \
             c not in [wage_rate_col, tot_lab_days_col, fam_lab_days_col]:
            labour_col = c

    # Handling merged/grouped columns:
    # If machinery and material mapped to same column (e.g. C4 had 'Machinery Material')
    if machinery_col and material_col and material_col == machinery_col:
        material_col = machinery_col + 1
        
    # Helper to find total in a total row
    def get_total_from_row(row_idx, preferred_col, line_idx=0):
        if not row_idx:
            return np.nan
        # 1. Preferred column
        if preferred_col:
            val = clean_num(sheet.cell(row_idx, preferred_col).value, line_idx=line_idx)
            if not np.isnan(val):
                return val
            # Adjacent columns
            for offset in [1, -1, 2]:
                c_try = preferred_col + offset
                if 1 <= c_try <= 15:
                    val = clean_num(sheet.cell(row_idx, c_try).value, line_idx=line_idx)
                    if not np.isnan(val) and val > 100:
                        return val
        # 2. Rightmost numeric value in row
        for c in range(15, 1, -1):
            val = clean_num(sheet.cell(row_idx, c).value, line_idx=line_idx)
            if not np.isnan(val):
                return val
        return np.nan

    # Check if total_incl_row contains 4 values that sum to the 4th (Labour + Machinery + Material = Total)
    incl_nums = [clean_num(sheet.cell(total_incl_row, c).value, line_idx=0) for c in range(2, 15)]
    incl_nums = [n for n in incl_nums if not np.isnan(n)]
    if len(incl_nums) == 4 and abs(incl_nums[0] + incl_nums[1] + incl_nums[2] - incl_nums[3]) <= 10:
        labour_cost = incl_nums[0]
        machinery_cost = incl_nums[1]
        material_cost = incl_nums[2]
        total_incl = incl_nums[3]
    else:
        # Extract raw cell numbers
        labour_cost = clean_num(sheet.cell(total_incl_row, labour_col).value, line_idx=0) if labour_col else np.nan
        machinery_cost = clean_num(sheet.cell(total_incl_row, machinery_col).value, line_idx=0) if machinery_col else np.nan
        material_cost = clean_num(sheet.cell(total_incl_row, material_col).value, line_idx=0) if material_col else np.nan
        total_incl = get_total_from_row(total_incl_row, total_col, line_idx=0)
    
    # Excl cost row
    if total_excl_row:
        total_excl_line_idx = 1 if (total_excl_row == total_incl_row) else 0
        excl_nums = [clean_num(sheet.cell(total_excl_row, c).value, line_idx=total_excl_line_idx) for c in range(2, 15)]
        excl_nums = [n for n in excl_nums if not np.isnan(n)]
        if len(excl_nums) == 4 and abs(excl_nums[0] + excl_nums[1] + excl_nums[2] - excl_nums[3]) <= 10:
            total_excl = excl_nums[3]
        else:
            total_excl = get_total_from_row(total_excl_row, total_col, line_idx=total_excl_line_idx)
    else:
        total_excl = np.nan
    
    # Mathematical fallback if total was empty or smaller than its components
    comps = [c for c in [labour_cost, machinery_cost, material_cost] if not np.isnan(c)]
    if comps:
        max_comp = max(comps)
        if np.isnan(total_incl) or total_incl < max_comp:
            total_incl = sum(comps)

    # Operation-level fallbacks for labor days/rate
    op_wage_rate = clean_num(sheet.cell(total_incl_row, wage_rate_col).value) if wage_rate_col else np.nan
    op_tot_days = clean_num(sheet.cell(total_incl_row, tot_lab_days_col).value) if tot_lab_days_col else np.nan
    op_fam_days = clean_num(sheet.cell(total_incl_row, fam_lab_days_col).value) if fam_lab_days_col else np.nan
    op_hired_days = np.nan
    if total_excl_row and tot_lab_days_col:
        op_hired_days = clean_num(sheet.cell(total_excl_row, tot_lab_days_col).value)

    res = {
        "Labour_Cost": labour_cost,
        "Machinery_Cost": machinery_cost,
        "Material_Cost": material_cost,
        "Total_Cost_Incl_Imputed": total_incl,
        "Total_Cost_Excl_Imputed": total_excl,
        "_op_wage_rate": op_wage_rate,
        "_op_tot_days": op_tot_days,
        "_op_fam_days": op_fam_days,
        "_op_hired_days": op_hired_days,
    }
    return res, op_row, total_incl_row, total_excl_row

# ==============================================================================
# 6. INPUTS BLOCK PARSER (Seed, Labour, Fertilizer, Remarks)
# ==============================================================================

def parse_inputs_block(sheet, start_row, end_row):
    """
    Extract Seed, Labour, Fertilizer quantities/prices, and Remarks.
    """
    data = {
        "Seed_Unit": "",
        "Seed_Quantity": np.nan,
        "Seed_Unit_Price": np.nan,
        "Labour_Wage_Rate": np.nan,
        "Hired_Labour_Days": np.nan,
        "Family_Labour_Days": np.nan,
        "Total_Labour_Days": np.nan,
        "Fertilizer_Urea_Price": np.nan,
        "Fertilizer_Urea_Qty": np.nan,
        "Fertilizer_TSP_Price": np.nan,
        "Fertilizer_TSP_Qty": np.nan,
        "Fertilizer_MOP_Price": np.nan,
        "Fertilizer_MOP_Qty": np.nan,
        "Fertilizer_Basal_TSP_Price": np.nan,
        "Fertilizer_Basal_TSP_Qty": np.nan,
        "Fertilizer_Basal_Urea_Price": np.nan,
        "Fertilizer_Basal_Urea_Qty": np.nan,
        "Fertilizer_Basal_MOP_Price": np.nan,
        "Fertilizer_Basal_MOP_Qty": np.nan,
        "Fertilizer_Top_Urea_Price": np.nan,
        "Fertilizer_Top_Urea_Qty": np.nan,
        "Fertilizer_Top_MOP_Price": np.nan,
        "Fertilizer_Top_MOP_Qty": np.nan,
        "Remarks": "",
    }
    
    unit_col = None
    qty_col = None
    price_col = None
    
    for r in range(start_row, min(start_row + 8, end_row)):
        row_str = " ".join([str(sheet.cell(r, c).value or "").strip().lower() for c in range(1, 10)])
        if "input" in row_str and ("unit" in row_str or "quantity" in row_str or "price" in row_str):
            for c in range(1, 10):
                cval = str(sheet.cell(r, c).value or "").strip().lower()
                if "unit price" in cval or "price" in cval:
                    price_col = c
                elif "unit" in cval and "price" not in cval:
                    unit_col = c
                elif "quant" in cval or "qty" in cval:
                    qty_col = c
            if unit_col and qty_col and price_col:
                break
                
    curr_section = None
    is_basal = False
    is_top = False
    remarks_list = []
    
    for r in range(start_row, end_row + 1):
        c1_val = str(sheet.cell(r, 1).value or "").strip()
        c1_lower = c1_val.lower().strip()
        row_str = " ".join([str(sheet.cell(r, c).value or "").strip().lower() for c in range(1, 12)])
        
        # Section header detection
        if c1_lower in ['seed', 'seeds', 'planting material', 'planting materials']:
            curr_section = "SEED"
            continue
        elif c1_lower in ['labour', 'labor']:
            curr_section = "LABOUR"
            continue
        elif 'fertilizer' in c1_lower and 'application' not in c1_lower and 'cost' not in c1_lower:
            curr_section = "FERTILIZER"
            continue
        elif "yield and return" in row_str:
            curr_section = "YIELD"
            continue
            
        # Remarks check
        if "remark" in row_str or "note" in row_str or c1_val.startswith("*"):
            full_r_txt = " ".join([str(sheet.cell(r, c).value or "").strip() for c in range(1, 10) if sheet.cell(r, c).value is not None])
            if full_r_txt:
                remarks_list.append(full_r_txt)
            continue
            
        if curr_section == "SEED":
            if np.isnan(data["Seed_Quantity"]) and np.isnan(data["Seed_Unit_Price"]):
                u = sheet.cell(r, unit_col).value if unit_col else None
                q = clean_num(sheet.cell(r, qty_col).value) if qty_col else np.nan
                p = clean_num(sheet.cell(r, price_col).value) if price_col else np.nan
                if np.isnan(q) or np.isnan(p):
                    nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 10)]
                    nums = [n for n in nums if not np.isnan(n)]
                    if len(nums) >= 2:
                        q, p = nums[0], nums[1]
                    elif len(nums) == 1:
                        p = nums[0]
                if u:
                    data["Seed_Unit"] = str(u).strip()
                elif sheet.cell(r, 2).value in ['kg', 'g', 'Bags', 'no', 'ac']:
                    data["Seed_Unit"] = str(sheet.cell(r, 2).value).strip()
                data["Seed_Quantity"] = q
                data["Seed_Unit_Price"] = p
                
        elif curr_section == "LABOUR":
            if "hired" in row_str:
                q = clean_num(sheet.cell(r, qty_col).value) if qty_col else np.nan
                p = clean_num(sheet.cell(r, price_col).value) if price_col else np.nan
                if np.isnan(q) or np.isnan(p):
                    nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 10)]
                    nums = [n for n in nums if not np.isnan(n)]
                    if len(nums) >= 2:
                        q, p = nums[0], nums[1]
                    elif len(nums) == 1:
                        if nums[0] > 100:
                            p = nums[0]
                        else:
                            q = nums[0]
                data["Hired_Labour_Days"] = q
                data["Labour_Wage_Rate"] = p
            elif "family" in row_str:
                q = clean_num(sheet.cell(r, qty_col).value) if qty_col else np.nan
                if np.isnan(q):
                    nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 10)]
                    nums = [n for n in nums if not np.isnan(n)]
                    if nums:
                        q = nums[0]
                data["Family_Labour_Days"] = q
            elif "total" in row_str and np.isnan(data["Total_Labour_Days"]):
                nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 10)]
                nums = [n for n in nums if not np.isnan(n)]
                if nums:
                    data["Total_Labour_Days"] = nums[0]
                    
        elif curr_section == "FERTILIZER":
            if "basal" in row_str:
                is_basal = True
                is_top = False
            elif "top" in row_str:
                is_basal = False
                is_top = True
                
            q = clean_num(sheet.cell(r, qty_col).value) if qty_col else np.nan
            p = clean_num(sheet.cell(r, price_col).value) if price_col else np.nan
            if np.isnan(q) or np.isnan(p):
                nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 10)]
                nums = [n for n in nums if not np.isnan(n)]
                if len(nums) == 3:
                    q, p = nums[1], nums[2]
                elif len(nums) == 2:
                    q, p = nums[0], nums[1]
                elif len(nums) == 1:
                    p = nums[0]
                    
            if "tsp" in row_str:
                data["Fertilizer_TSP_Price"] = p
                data["Fertilizer_TSP_Qty"] = q
                if is_basal:
                    data["Fertilizer_Basal_TSP_Price"] = p
                    data["Fertilizer_Basal_TSP_Qty"] = q
            elif "urea" in row_str:
                data["Fertilizer_Urea_Price"] = p
                data["Fertilizer_Urea_Qty"] = q
                if is_basal:
                    data["Fertilizer_Basal_Urea_Price"] = p
                    data["Fertilizer_Basal_Urea_Qty"] = q
                elif is_top:
                    data["Fertilizer_Top_Urea_Price"] = p
                    data["Fertilizer_Top_Urea_Qty"] = q
            elif "mop" in row_str:
                data["Fertilizer_MOP_Price"] = p
                data["Fertilizer_MOP_Qty"] = q
                if is_basal:
                    data["Fertilizer_Basal_MOP_Price"] = p
                    data["Fertilizer_Basal_MOP_Qty"] = q
                elif is_top:
                    data["Fertilizer_Top_MOP_Price"] = p
                    data["Fertilizer_Top_MOP_Qty"] = q

    if remarks_list:
        data["Remarks"] = " | ".join(remarks_list)
        
    return data

# ==============================================================================
# 7. YIELD AND RETURNS BLOCK PARSER
# ==============================================================================

def parse_yield_block(sheet, start_row, max_row):
    """
    Extract Yield and Returns per acre figures:
    Average Yield (kg/ac), Price of Produce (Rs/kg), Gross Income (Rs/ac),
    Profit Incl/Excl Imputed Cost (Rs/ac), Per Unit Cost Incl/Excl Imputed (Rs/kg).
    Guarded against stray secondary unedited template blocks and shifted layouts.
    """
    data = {
        "Average_Yield_kg_ac": np.nan,
        "Price_of_Produce": np.nan,
        "Gross_Income": np.nan,
        "Profit_Incl_Imputed": np.nan,
        "Profit_Excl_Imputed": np.nan,
        "Per_Unit_Cost_Incl_Imputed": np.nan,
        "Per_Unit_Cost_Excl_Imputed": np.nan,
    }
    
    remarks = []
    
    # Check if header row itself has Price (shifted layout like Black gram-Vavuniya)
    hdr_shifted_price = None
    for r in range(start_row, min(start_row + 5, max_row + 1)):
        r_str = " ".join([str(sheet.cell(r, c).value or "").strip().lower() for c in range(1, 12)])
        if "yield and return" in r_str:
            nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 12)]
            nums = [n for n in nums if not np.isnan(n)]
            if nums and nums[0] > 0:
                hdr_shifted_price = nums[0]
            break
            
    for r in range(start_row, max_row + 1):
        row_str = " ".join([str(sheet.cell(r, c).value or "").strip().lower() for c in range(1, 12)])
        nums = [clean_num(sheet.cell(r, c).value) for c in range(2, 12)]
        nums = [n for n in nums if not np.isnan(n)]
        
        c1_str = str(sheet.cell(r, 1).value or "").strip()
        if c1_str.startswith("*") or "remark" in row_str or "note" in row_str:
            full_txt = " ".join([str(sheet.cell(r, c).value or "").strip() for c in range(1, 10) if sheet.cell(r, c).value is not None])
            if full_txt:
                remarks.append(full_txt)
                
        # If unedited template error row encountered, stop parsing further
        if '#value!' in row_str:
            break
            
        if "yield" in row_str and "average" in row_str:
            if nums and np.isnan(data["Average_Yield_kg_ac"]):
                data["Average_Yield_kg_ac"] = nums[0]
        elif "price of produce" in row_str:
            if nums and np.isnan(data["Price_of_Produce"]):
                # Check for shifted layout
                if hdr_shifted_price is not None and nums[0] > 1000 and not np.isnan(data["Average_Yield_kg_ac"]) and data["Average_Yield_kg_ac"] < 1000:
                    data["Price_of_Produce"] = hdr_shifted_price
                    data["Gross_Income"] = nums[0]
                else:
                    data["Price_of_Produce"] = nums[0]
        elif "gross income" in row_str:
            if nums and np.isnan(data["Gross_Income"]):
                data["Gross_Income"] = nums[0]
            elif nums and hdr_shifted_price is not None and np.isnan(data["Profit_Incl_Imputed"]):
                data["Profit_Incl_Imputed"] = nums[0]
        elif "profit" in row_str and ("incl" in row_str or "including" in row_str):
            if nums and np.isnan(data["Profit_Incl_Imputed"]):
                data["Profit_Incl_Imputed"] = nums[0]
            elif nums and hdr_shifted_price is not None and np.isnan(data["Profit_Excl_Imputed"]):
                data["Profit_Excl_Imputed"] = nums[0]
        elif "profit" in row_str and ("excl" in row_str or "excluding" in row_str):
            if nums and np.isnan(data["Profit_Excl_Imputed"]):
                data["Profit_Excl_Imputed"] = nums[0]
            elif nums and hdr_shifted_price is not None and np.isnan(data["Per_Unit_Cost_Incl_Imputed"]):
                data["Per_Unit_Cost_Incl_Imputed"] = nums[0]
        elif "unit cost" in row_str and ("incl" in row_str or "including" in row_str):
            if nums and np.isnan(data["Per_Unit_Cost_Incl_Imputed"]):
                data["Per_Unit_Cost_Incl_Imputed"] = nums[0]
            elif nums and hdr_shifted_price is not None and np.isnan(data["Per_Unit_Cost_Excl_Imputed"]):
                data["Per_Unit_Cost_Excl_Imputed"] = nums[0]
        elif "unit cost" in row_str and ("excl" in row_str or "excluding" in row_str):
            if nums and np.isnan(data["Per_Unit_Cost_Excl_Imputed"]):
                data["Per_Unit_Cost_Excl_Imputed"] = nums[0]
                
    if remarks:
        data["_yield_remarks"] = " | ".join(remarks)
    else:
        data["_yield_remarks"] = ""
        
    return data

# ==============================================================================
# 8. MASTER SHEET PARSER
# ==============================================================================

def parse_sheet(wb_path, sheet, sname, folder_info):
    """
    Parse a single cultivation sheet and return a complete record dictionary.
    """
    crop, district, irrigation = parse_sheet_metadata(wb_path, sheet, sname)
    
    # Parse Operation Table
    op_data, op_row, incl_row, excl_row = parse_operation_table(sheet)
    
    # Determine bounds for inputs & yield blocks
    inputs_start = (excl_row or incl_row or 20) + 1
    max_r = sheet.max_row
    
    # Find yield start row
    yield_start = inputs_start
    for r in range(inputs_start, min(max_r + 1, inputs_start + 35)):
        r_str = " ".join([str(sheet.cell(r, c).value or "").strip().lower() for c in range(1, 10)])
        if "yield and return" in r_str:
            yield_start = r
            break
            
    # Parse inputs block (from inputs_start up to yield_start)
    inputs_data = parse_inputs_block(sheet, inputs_start, yield_start)
    
    # Parse yield block (from yield_start to sheet end)
    yield_data = parse_yield_block(sheet, yield_start, max_r)
    
    # Merge remarks
    remarks_combined = []
    if inputs_data.get("Remarks"):
        remarks_combined.append(inputs_data["Remarks"])
    if yield_data.get("_yield_remarks"):
        remarks_combined.append(yield_data["_yield_remarks"])
    final_remarks = " | ".join(remarks_combined) if remarks_combined else ""
    
    # Labour days/rate fallback from operations table if inputs block was missing them
    hired_days = inputs_data.get("Hired_Labour_Days")
    if np.isnan(hired_days) or hired_days is None:
        hired_days = op_data.get("_op_hired_days", np.nan)
        
    family_days = inputs_data.get("Family_Labour_Days")
    if np.isnan(family_days) or family_days is None:
        family_days = op_data.get("_op_fam_days", np.nan)
        
    tot_days = inputs_data.get("Total_Labour_Days")
    if np.isnan(tot_days) or tot_days is None:
        tot_days = op_data.get("_op_tot_days", np.nan)
        
    wage_rate = inputs_data.get("Labour_Wage_Rate")
    if np.isnan(wage_rate) or wage_rate is None:
        wage_rate = op_data.get("_op_wage_rate", np.nan)
        
    record = {
        "Season": folder_info["Season"],
        "Year_Start": folder_info["Year_Start"],
        "Year_End": folder_info["Year_End"],
        "Year_Range": folder_info["Year_Range"],
        "Crop": crop,
        "District": district,
        "Irrigation_Type": irrigation,
        "Field_Size_Basis": "Per Acre",
        "Labour_Cost": op_data.get("Labour_Cost", np.nan),
        "Machinery_Cost": op_data.get("Machinery_Cost", np.nan),
        "Material_Cost": op_data.get("Material_Cost", np.nan),
        "Total_Cost_Incl_Imputed": op_data.get("Total_Cost_Incl_Imputed", np.nan),
        "Total_Cost_Excl_Imputed": op_data.get("Total_Cost_Excl_Imputed", np.nan),
        "Seed_Unit": inputs_data.get("Seed_Unit", ""),
        "Seed_Quantity": inputs_data.get("Seed_Quantity", np.nan),
        "Seed_Unit_Price": inputs_data.get("Seed_Unit_Price", np.nan),
        "Labour_Wage_Rate": wage_rate,
        "Hired_Labour_Days": hired_days,
        "Family_Labour_Days": family_days,
        "Total_Labour_Days": tot_days,
        "Fertilizer_Urea_Price": inputs_data.get("Fertilizer_Urea_Price", np.nan),
        "Fertilizer_Urea_Qty": inputs_data.get("Fertilizer_Urea_Qty", np.nan),
        "Fertilizer_TSP_Price": inputs_data.get("Fertilizer_TSP_Price", np.nan),
        "Fertilizer_TSP_Qty": inputs_data.get("Fertilizer_TSP_Qty", np.nan),
        "Fertilizer_MOP_Price": inputs_data.get("Fertilizer_MOP_Price", np.nan),
        "Fertilizer_MOP_Qty": inputs_data.get("Fertilizer_MOP_Qty", np.nan),
        "Fertilizer_Basal_TSP_Price": inputs_data.get("Fertilizer_Basal_TSP_Price", np.nan),
        "Fertilizer_Basal_TSP_Qty": inputs_data.get("Fertilizer_Basal_TSP_Qty", np.nan),
        "Fertilizer_Basal_Urea_Price": inputs_data.get("Fertilizer_Basal_Urea_Price", np.nan),
        "Fertilizer_Basal_Urea_Qty": inputs_data.get("Fertilizer_Basal_Urea_Qty", np.nan),
        "Fertilizer_Basal_MOP_Price": inputs_data.get("Fertilizer_Basal_MOP_Price", np.nan),
        "Fertilizer_Basal_MOP_Qty": inputs_data.get("Fertilizer_Basal_MOP_Qty", np.nan),
        "Fertilizer_Top_Urea_Price": inputs_data.get("Fertilizer_Top_Urea_Price", np.nan),
        "Fertilizer_Top_Urea_Qty": inputs_data.get("Fertilizer_Top_Urea_Qty", np.nan),
        "Fertilizer_Top_MOP_Price": inputs_data.get("Fertilizer_Top_MOP_Price", np.nan),
        "Fertilizer_Top_MOP_Qty": inputs_data.get("Fertilizer_Top_MOP_Qty", np.nan),
        "Average_Yield_kg_ac": yield_data.get("Average_Yield_kg_ac", np.nan),
        "Price_of_Produce": yield_data.get("Price_of_Produce", np.nan),
        "Gross_Income": yield_data.get("Gross_Income", np.nan),
        "Profit_Incl_Imputed": yield_data.get("Profit_Incl_Imputed", np.nan),
        "Profit_Excl_Imputed": yield_data.get("Profit_Excl_Imputed", np.nan),
        "Per_Unit_Cost_Incl_Imputed": yield_data.get("Per_Unit_Cost_Incl_Imputed", np.nan),
        "Per_Unit_Cost_Excl_Imputed": yield_data.get("Per_Unit_Cost_Excl_Imputed", np.nan),
        "Remarks": final_remarks,
        "Source_File": wb_path.name,
        "Source_Sheet": sname,
    }
    
    return record

# ==============================================================================
# 9. MAIN PIPELINE EXECUTION
# ==============================================================================

def run_consolidation(workspace_dir):
    workspace = Path(workspace_dir)
    surveys_dir = workspace / "data" / "raw_surveys"
    if not surveys_dir.exists():
        surveys_dir = workspace
    season_folders = sorted([f for f in surveys_dir.iterdir() if f.is_dir() and not f.name.startswith('.')])
    
    all_records = []
    parse_reports = []
    
    print(f"Starting consolidation across {len(season_folders)} season folders...")
    
    for fld in season_folders:
        folder_info = parse_season_folder(fld.name)
        xlsx_files = sorted(list(fld.rglob("*.xlsx")))
        print(f"\nProcessing Folder: {fld.name} ({len(xlsx_files)} workbooks)")
        
        folder_record_count = 0
        
        for xf in xlsx_files:
            rel_file = xf.relative_to(workspace)
            try:
                wb = openpyxl.load_workbook(xf, data_only=True)
            except Exception as e:
                parse_reports.append({
                    "Folder": fld.name,
                    "File": xf.name,
                    "Sheet": "ALL",
                    "Status": "FAILED_OPEN",
                    "Reason": f"Could not load workbook: {e}",
                    "Crop_Detected": "",
                    "District_Detected": "",
                })
                continue
                
            prev_record_idx = None
            
            for sname in wb.sheetnames:
                sheet = wb[sname]
                cls_type, cls_reason = classify_sheet(sheet, sname)
                
                if cls_type in ["BLANK", "COVER_TOC", "SUMMARY", "DRAFT"]:
                    parse_reports.append({
                        "Folder": fld.name,
                        "File": xf.name,
                        "Sheet": sname,
                        "Status": f"SKIPPED_{cls_type}",
                        "Reason": cls_reason,
                        "Crop_Detected": "",
                        "District_Detected": "",
                    })
                    continue
                    
                if cls_type == "CONTINUATION":
                    # Merge yield data into the previous record if available
                    if prev_record_idx is not None and prev_record_idx < len(all_records):
                        cont_yield = parse_yield_block(sheet, 1, sheet.max_row)
                        for yk in ["Average_Yield_kg_ac", "Price_of_Produce", "Gross_Income",
                                   "Profit_Incl_Imputed", "Profit_Excl_Imputed",
                                   "Per_Unit_Cost_Incl_Imputed", "Per_Unit_Cost_Excl_Imputed"]:
                            if not np.isnan(cont_yield.get(yk, np.nan)):
                                all_records[prev_record_idx][yk] = cont_yield[yk]
                        if cont_yield.get("_yield_remarks"):
                            old_rem = all_records[prev_record_idx]["Remarks"]
                            all_records[prev_record_idx]["Remarks"] = (old_rem + " | " + cont_yield["_yield_remarks"]).strip(" |")
                            
                        parse_reports.append({
                            "Folder": fld.name,
                            "File": xf.name,
                            "Sheet": sname,
                            "Status": "MERGED_CONTINUATION",
                            "Reason": f"Merged yield/returns into preceding sheet: {all_records[prev_record_idx]['Source_Sheet']}",
                            "Crop_Detected": all_records[prev_record_idx]["Crop"],
                            "District_Detected": all_records[prev_record_idx]["District"],
                        })
                    else:
                        parse_reports.append({
                            "Folder": fld.name,
                            "File": xf.name,
                            "Sheet": sname,
                            "Status": "SKIPPED_ORPHAN_CONTINUATION",
                            "Reason": "Yield block continuation sheet with no preceding cultivation sheet to attach to",
                            "Crop_Detected": "",
                            "District_Detected": "",
                        })
                    continue
                    
                # Full COC sheet parsing
                record = parse_sheet(xf, sheet, sname, folder_info)
                
                # Check for identical duplicate report within the same season
                is_duplicate = False
                dup_match_info = ""
                for existing in all_records:
                    if existing["Season"] == record["Season"] and \
                       existing["Year_Range"] == record["Year_Range"] and \
                       existing["Crop"] == record["Crop"] and \
                       existing["District"] == record["District"] and \
                       existing["Irrigation_Type"] == record["Irrigation_Type"]:
                        # Compare core figures
                        num_keys = ["Labour_Cost", "Machinery_Cost", "Material_Cost", 
                                    "Total_Cost_Incl_Imputed", "Total_Cost_Excl_Imputed",
                                    "Average_Yield_kg_ac", "Gross_Income"]
                        matches = 0
                        total_valid = 0
                        for k in num_keys:
                            v1 = existing.get(k, np.nan)
                            v2 = record.get(k, np.nan)
                            if not np.isnan(v1) and not np.isnan(v2):
                                total_valid += 1
                                if abs(v1 - v2) <= 1.0 or (abs(v1 - v2) / max(abs(v1), abs(v2)) < 0.001):
                                    matches += 1
                        if total_valid >= 3 and matches == total_valid:
                            is_duplicate = True
                            dup_match_info = f"Identical report already extracted from {existing['Source_File']} / sheet '{existing['Source_Sheet']}'"
                            break
                            
                if is_duplicate:
                    parse_reports.append({
                        "Folder": fld.name,
                        "File": xf.name,
                        "Sheet": sname,
                        "Status": "DROPPED_IDENTICAL_DUPLICATE",
                        "Reason": dup_match_info,
                        "Crop_Detected": record["Crop"] or "",
                        "District_Detected": record["District"] or "",
                    })
                    continue

                # Check for critical missing values
                missing_fields = []
                if not record["Crop"]:
                    missing_fields.append("Crop")
                if not record["District"]:
                    missing_fields.append("District")
                if np.isnan(record["Total_Cost_Incl_Imputed"]):
                    missing_fields.append("Total_Cost_Incl_Imputed")
                    
                if missing_fields:
                    status = "PARTIALLY_PARSED"
                    reason = f"Missing key fields: {', '.join(missing_fields)}"
                else:
                    status = "SUCCESS"
                    reason = "Successfully extracted cost, yield, and input parameters"
                    
                all_records.append(record)
                prev_record_idx = len(all_records) - 1
                folder_record_count += 1
                
                parse_reports.append({
                    "Folder": fld.name,
                    "File": xf.name,
                    "Sheet": sname,
                    "Status": status,
                    "Reason": reason,
                    "Crop_Detected": record["Crop"] or "",
                    "District_Detected": record["District"] or "",
                })
                
            wb.close()
            
        print(f"  -> Extracted {folder_record_count} valid records from {fld.name}")
        
    # Convert to DataFrames
    df_master = pd.DataFrame(all_records)
    df_report = pd.DataFrame(parse_reports)
    
    # Save master CSV
    output_dir = workspace / "data" / "processed" if (workspace / "data" / "processed").exists() else workspace
    master_csv_path = output_dir / "coc_master.csv"
    report_csv_path = output_dir / "parse_report.csv"
    
    df_master.to_csv(master_csv_path, index=False, encoding="utf-8-sig")
    df_report.to_csv(report_csv_path, index=False, encoding="utf-8-sig")
    
    print("\n" + "=" * 80)
    print("CONSOLIDATION COMPLETE!")
    print(f"Master CSV saved to: {master_csv_path}")
    print(f"Parse Report saved to: {report_csv_path}")
    print("=" * 80)
    
    # Print Sanity Checks
    print("\n--- SANITY CHECK SUMMARY ---")
    print(f"Total Rows Extracted: {len(df_master)}")
    print(f"Total Sheets Audited in Parse Report: {len(df_report)}")
    
    print("\n1. Record Count by Season Folder:")
    folder_summary = df_report[df_report["Status"].isin(["SUCCESS", "PARTIALLY_PARSED"])].groupby("Folder").size()
    for fld in season_folders:
        cnt = folder_summary.get(fld.name, 0)
        status_flag = "OK" if cnt > 0 else "WARNING: ZERO ROWS EXTRACTED!"
        print(f"   {fld.name:20s}: {cnt:3d} rows  [{status_flag}]")
        
    print(f"\n2. Distinct Crops ({df_master['Crop'].nunique()}):")
    sorted_crops = sorted([str(c) for c in df_master['Crop'].unique() if pd.notna(c)])
    print(", ".join(sorted_crops))
    
    print(f"\n3. Distinct Districts ({df_master['District'].nunique()}):")
    sorted_districts = sorted([str(d) for d in df_master['District'].unique() if pd.notna(d)])
    print(", ".join(sorted_districts))
    
    print("\n4. Missing / Null Counts per Column:")
    null_counts = df_master.isnull().sum()
    for col, n_cnt in null_counts.items():
        pct = (n_cnt / len(df_master)) * 100
        print(f"   {col:30s}: {n_cnt:4d} nulls ({pct:5.1f}%)")
        
    print("\n5. Parse Report Status Summary:")
    print(df_report["Status"].value_counts().to_string())
    
    # Check duplicate combinations in final dataset
    group_cols = ["Season", "Year_Range", "Crop", "District", "Irrigation_Type"]
    remaining_dups = df_master[df_master.duplicated(subset=group_cols, keep=False)]
    print(f"\n6. Remaining Duplicate Groups in Final Dataset: {len(remaining_dups)}")

if __name__ == "__main__":
    ws_path = r"c:\Users\MSI\Desktop\Research\COC folder"
    if len(sys.argv) > 1:
        ws_path = sys.argv[1]
    run_consolidation(ws_path)
